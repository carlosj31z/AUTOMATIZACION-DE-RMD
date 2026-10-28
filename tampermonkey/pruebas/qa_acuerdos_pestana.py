"""Pruebas de la pestaña «Acuerdos» (v1.103) sobre la misma Configuración Maestra con OpenUI5 real de qa_controles_cambio.py (ver allí cómo preparar
RMD_UI5_RES). Supabase de status-rmd se simula en memoria: no escribe nada real.
    RMD_UI5_RES=/ruta/a/res python tampermonkey/pruebas/qa_acuerdos_pestana.py
"""
import os, sys, pathlib, json, tempfile, base64, re
from urllib.parse import urlparse, parse_qs, unquote
from playwright.sync_api import sync_playwright
RES = pathlib.Path(os.environ.get('RMD_UI5_RES', 'ui5/res'))
PAGE = r"""<!doctype html><html><head><meta charset="utf-8"><title>maestra</title>
<script id="sap-ui-bootstrap" src="/resources/sap-ui-core.js" data-sap-ui-theme="sap_fiori_3" data-sap-ui-libs="sap.m,sap.f" data-sap-ui-async="true" data-sap-ui-compatVersion="edge"></script>
<script>
sap.ui.getCore().attachInit(function () {
  sap.ui.require(['sap/m/Dialog','sap/m/IconTabBar','sap/m/IconTabFilter','sap/m/IconTabSeparator','sap/m/Table','sap/m/Column','sap/m/Text','sap/m/Button','sap/m/ColumnListItem','sap/m/Label','sap/m/Input','sap/m/VBox','sap/f/DynamicPage','sap/f/DynamicPageHeader'], function (Dialog, ITB, ITF, ITS, Table, Column, Text, Button, CLI, Label, Input, VBox, DP, DPH) {
    var tablas = {};
    var mk = function (k, t, ic) { var tb = new Table({ headerText: t, visible: k === '1', columns: [new Column({ header: new Text({ text: 'Código' }) }), new Column({ header: new Text({ text: 'Descripción' }) })] });
      for (var i = 1; i < 40; i++) tb.addItem(new CLI({ cells: [new Text({ text: String(i) }), new Text({ text: t + ' ' + i })] }));
      tablas[k] = tb; return new ITF({ key: k, text: t, icon: ic }); };
    var itb = new ITB('itbMaestra', { expandable: false, select: function (e) { var k = e.getParameter('key'); window.__sel = k; Object.keys(tablas).forEach(function (x) { tablas[x].setVisible(x === k); }); }, items: [mk('1', 'Estructura', 'sap-icon://building'), new ITS(), mk('2', 'Etiqueta', 'sap-icon://tag'), new ITS(), mk('3', 'Paso', 'sap-icon://process'), new ITS(), mk('4', 'Motivo', 'sap-icon://undo'), new ITS(), mk('5', 'Utensilios', 'sap-icon://laptop'), new ITS(), mk('6', 'Motivo de lapsos', 'sap-icon://journey-arrive')] });
    var cuerpo = new VBox({ items: Object.keys(tablas).map(function (k) { return tablas[k]; }) });
    var dp = new DP('frgConfiguration--dynamicPageId', { header: new DPH({ pinnable: false, content: [itb, new Label({ text: 'Descripción:' }), new Input({ placeholder: 'Ingrese Descripción' })] }), content: cuerpo });
    var d = new Dialog({ title: 'Configuracion Maestra', contentWidth: '95%', contentHeight: '90%', verticalScrolling: false, content: [dp], endButton: new Button({ text: 'Cerrar', press: function () { d.close(); } }) });
    window.__abrir = function () { d.open(); }; d.open(); window.__listo = true;
});
});
</script></head><body class="sapUiBody"></body></html>
"""
SCRIPT = (pathlib.Path(__file__).resolve().parents[1] / 'rmd-ui-mejoras.user.js').read_text(encoding='utf-8')
OUT = pathlib.Path(tempfile.gettempdir()) / 'qa_acuerdos_pestana'; OUT.mkdir(exist_ok=True)
CHROME = os.environ.get('RMD_CHROME') or None
# PNG de 40x30 (captura de ejemplo)
PNG = OUT / 'captura.png'
def png(w, h):
    import zlib, struct
    fila = lambda y: b'\x00' + b''.join(bytes((x * 6 % 256, y * 8 % 256, 160)) for x in range(w))
    cr = lambda t, d: struct.pack('>I', len(d)) + t + d + struct.pack('>I', zlib.crc32(t + d) & 0xffffffff)
    return b'\x89PNG\r\n\x1a\n' + cr(b'IHDR', struct.pack('>IIBBBBB', w, h, 8, 2, 0, 0, 0)) + cr(b'IDAT', zlib.compress(b''.join(fila(y) for y in range(h)))) + cr(b'IEND', b'')
PNG.write_bytes(png(40, 30))
kv = {}
# fila escrita antes por status-rmd, con campos que el portal no usa
kv['acuerdo:sr1'] = {'id': 'sr1', 'titulo': 'Precauciones', 'planta': 'ambas', 'etapas': [], 'seccion': 'PRECAUCIONES', 'tipo': 'agregar', 'pasos': [{'codigo': '180998', 'descripcion': 'CONTROLAR'}],
                     'creadoPor': 'Carlos', 'creadoEn': '2026-10-28T10:00:00Z', 'actualizadoEn': '2026-10-28T10:00:00Z', 'historial': [{'fecha': '2026-10-28T10:00:00Z', 'por': 'Carlos', 'accion': 'Creado'}], 'soloStatusRmd': 7}
res = []
def ok(n, c, d=''): res.append((n, bool(c))); print(('PASA  ' if c else 'FALLA ') + n + (' — ' + str(d)[:300] if d and not c else ''), flush=True)
def ruta(route):
    req = route.request; u = req.url
    if 'supabase.co/rest/v1/rmd_dashboard_kv' in u:
        q = parse_qs(urlparse(u).query); key = (q.get('key') or [''])[0]
        if req.method == 'GET':
            pre = key[5:].rstrip('*') if key.startswith('like.') else None
            filas = [{'key': k, 'value': v} for k, v in kv.items() if pre is not None and k.startswith(pre)]
            return route.fulfill(body=json.dumps(filas), content_type='application/json', headers={'Access-Control-Allow-Origin': '*'})
        if req.method == 'POST':
            for f in json.loads(req.post_data): kv[f['key']] = f['value']
            return route.fulfill(status=201, body='', headers={'Access-Control-Allow-Origin': '*'})
        if req.method == 'DELETE':
            kv.pop(key[3:], None); return route.fulfill(status=204, body='', headers={'Access-Control-Allow-Origin': '*'})
        return route.fulfill(status=200, body='', headers={'Access-Control-Allow-Origin': '*', 'Access-Control-Allow-Headers': '*', 'Access-Control-Allow-Methods': '*'})
    if '/resources/' in u:
        p = RES / u.split('/resources/', 1)[1].split('?')[0]
        return route.fulfill(path=str(p)) if p.exists() else route.fulfill(status=404, body='')
    return route.fulfill(body=PAGE, content_type='text/html')
TABS = "[...document.querySelectorAll('[role=tab]')].map(t=>t.textContent.trim())"
VIS = "(t) => [...document.querySelectorAll('.sapMListTbl')].some((x) => x.getClientRects().length && (x.closest('.sapMList').textContent || '').includes(t + ' 1'))"
with sync_playwright() as pw:
    b = pw.chromium.launch(executable_path=CHROME); ctx = b.new_context(viewport={'width': 1400, 'height': 860}, accept_downloads=True); pg = ctx.new_page()
    errs = []; pg.on('pageerror', lambda e: errs.append(str(e)))
    pg.route('**/*', ruta); pg.goto('https://x.hana.ondemand.com/cp.portal/ui5appruntime.html')
    pg.wait_for_function('window.__listo', timeout=60000); pg.wait_for_timeout(800)
    pg.add_script_tag(content=SCRIPT); pg.wait_for_timeout(2000)
    tabs = pg.evaluate(TABS)
    ok('pestaña «Acuerdos» a la derecha de «Controles de Cambios»', tabs[-2].endswith('Controles de Cambios') and tabs[-1].endswith('Acuerdos'), tabs)
    dlg0 = pg.evaluate("(() => { const r = document.querySelector('.sapMDialog').getBoundingClientRect(); return [r.width, r.height]; })()")
    pg.locator('[role=tab]', has_text='Acuerdos').click(); pg.wait_for_timeout(1200)
    ok('contenido visible', pg.locator('.rmd-acu-tab').is_visible())
    ok('lee los acuerdos de status-rmd', pg.locator('.rmd-acu-item').count() == 1 and 'fuente: status-rmd' in pg.locator('.rmd-acu-info').inner_text(), pg.locator('.rmd-acu-info').inner_text())
    ok('sin la tabla del portal debajo', not pg.evaluate("(" + VIS + ")('Estructura')"))
    ok('sin los filtros del portal', not pg.evaluate("[...document.querySelectorAll('.sapFDynamicPageHeader input')].some((i) => !i.closest('.rmd-acu-tab') && i.getClientRects().length)"))
    ok('sin el botón de los Word', pg.locator('text=Word').count() == 0)
    pg.screenshot(path=str(OUT / '1-lista.png'))
    # nuevo acuerdo
    pg.locator('.rmd-acu-barra button', has_text='Nuevo acuerdo').click(); pg.wait_for_timeout(400)
    ed = pg.locator('.rmd-acu-editor'); ok('editor abierto', ed.is_visible())
    ed.locator('[data-c=titulo]').fill('Despejes sin V°B° en ATE')
    ed.locator('[data-c=planta]').select_option('planta1'); ed.locator('[data-etapa=FABRICACION]').check()
    ed.locator('[data-c=faseDesde]').select_option('2'); ed.locator('[data-c=tipo]').select_option('reemplazar')
    ed.locator('[data-c=detalle]').fill('Desde la fase 2 el despeje ya no lleva V°B°.\nSe cambia por el paso 184004.')
    ed.locator('button', has_text='Guardar acuerdo').click(); pg.wait_for_timeout(300)
    ok('pide un código si no es «Indicación»', 'necesita al menos un código' in ed.locator('.rmd-acu-error').inner_text())
    ed.locator('button', has_text='+ Paso de referencia').click()
    ed.locator('.rmd-acu-pasos tbody tr').last.locator('[data-p=codigo]').fill('184004'); ed.locator('.rmd-acu-pasos tbody tr').last.locator('[data-p=descripcion]').fill('REALIZAR EL DESPEJE')
    ed.locator('.rmd-acu-pasos tbody tr').last.locator('[data-p=rPor]').check()
    ed.locator('button', has_text='+ Paso de referencia').click()
    ed.locator('.rmd-acu-pasos tbody tr').last.locator('[data-p=codigo]').fill('18A')
    ed.locator('button', has_text='Guardar acuerdo').click(); pg.wait_for_timeout(300)
    ok('valida código no numérico', '18A' in ed.locator('.rmd-acu-error').inner_text())
    ed.locator('.rmd-acu-pasos tbody tr').last.locator('[data-a=quitar]').click()
    # imagen elegida + imagen pegada
    with pg.expect_file_chooser() as fc: ed.locator('button', has_text='Elegir imagen').click()
    fc.value.set_files(str(PNG)); pg.wait_for_timeout(800)
    pg.evaluate("""async (b64) => { const bin = Uint8Array.from(atob(b64), (c) => c.charCodeAt(0)), dt = new DataTransfer();
      dt.items.add(new File([bin], 'image.png', { type: 'image/png' })); document.querySelector('.rmd-acu-editor [data-c=detalle]').dispatchEvent(new ClipboardEvent('paste', { clipboardData: dt, bubbles: true })); }""", base64.b64encode(PNG.read_bytes()).decode())
    pg.wait_for_timeout(800)
    ok('2 imágenes en el editor (elegida y pegada)', ed.locator('.rmd-acu-imgs-ed img').count() == 2)
    ed.locator('[data-img-nombre="0"]').fill('Captura del portal')
    pg.screenshot(path=str(OUT / '2-editor.png'))
    ed.locator('button', has_text='Guardar acuerdo').click(); pg.wait_for_timeout(800)
    nuevos = [v for k, v in kv.items() if k != 'acuerdo:sr1']
    ok('guardado en Supabase (una fila acuerdo:<id>)', len(nuevos) == 1, list(kv))
    n = nuevos[0] if nuevos else {}
    ok('campos guardados', n.get('planta') == 'planta1' and n.get('etapas') == ['FABRICACION'] and n.get('faseDesde') == 2 and n.get('tipo') == 'reemplazar' and n['pasos'][0]['codigo'] == '184004' and n['pasos'][0]['rPor'], {k: n.get(k) for k in ('planta', 'etapas', 'faseDesde', 'tipo', 'pasos')})
    ok('imágenes comprimidas como data URL', len(n.get('imagenes', [])) == 2 and all(re.match(r'data:image/(webp|jpeg);base64,', i['src']) for i in n['imagenes']) and n['imagenes'][0]['nombre'] == 'Captura del portal', [(i.get('nombre'), i['src'][:30]) for i in n.get('imagenes', [])])
    ok('quién y historial', n.get('creadoPor') == 'Portal SAP' and len(n.get('historial', [])) == 1 and n['historial'][0]['accion'] == 'Creado desde el portal')
    ok('editor cerrado y lista con 2', pg.locator('.rmd-acu-editor').count() == 0 and pg.locator('.rmd-acu-item').count() == 2)
    it = pg.locator('.rmd-acu-item', has_text='Despejes sin V°B° en ATE')
    ok('tarjeta con miniaturas y texto en varias líneas', it.locator('.rmd-acu-imgs img').count() == 2 and 'Se cambia por el paso 184004' in it.inner_text())
    pg.screenshot(path=str(OUT / '3-lista.png'))
    it.locator('.rmd-acu-imgs img').first.click(); pg.wait_for_timeout(300)
    ok('visor de imagen', pg.locator('.rmd-acu-visor').is_visible()); pg.screenshot(path=str(OUT / '4-visor.png'))
    pg.keyboard.press('Escape'); pg.wait_for_timeout(200); ok('Esc cierra el visor y no la Configuración Maestra', pg.locator('.rmd-acu-visor').count() == 0 and pg.locator('.sapMDialog').is_visible())
    # filtros
    pg.locator('.rmd-acu-barra [data-f=planta]').select_option('planta2'); pg.wait_for_timeout(200)
    ok('filtro Planta 2 oculta el de ATE', pg.locator('.rmd-acu-item').count() == 1)
    pg.locator('.rmd-acu-barra [data-f=planta]').select_option(''); pg.locator('.rmd-acu-barra [data-f=q]').fill('184004'); pg.wait_for_timeout(200)
    ok('buscar por código', pg.locator('.rmd-acu-item').count() == 1); pg.locator('.rmd-acu-barra [data-f=q]').fill(''); pg.wait_for_timeout(200)
    # editar el de status-rmd: conserva sus campos
    pg.locator('.rmd-acu-item', has_text='Precauciones').locator('[data-acc=editar]').click(); pg.wait_for_timeout(300)
    pg.locator('.rmd-acu-editor [data-c=titulo]').fill('Precauciones (paso 180998)'); pg.locator('.rmd-acu-editor button', has_text='Guardar cambios').click(); pg.wait_for_timeout(700)
    s = kv['acuerdo:sr1']
    ok('editar conserva lo de status-rmd', s['titulo'] == 'Precauciones (paso 180998)' and s.get('soloStatusRmd') == 7 and s['creadoPor'] == 'Carlos' and len(s['historial']) == 2, {k: s.get(k) for k in ('titulo', 'soloStatusRmd', 'creadoPor')})
    # anular
    pg.locator('.rmd-acu-item', has_text='Precauciones').locator('[data-acc=estado]').click(); pg.wait_for_timeout(700)
    ok('anular', kv['acuerdo:sr1']['estado'] == 'anulado' and pg.locator('.rmd-acu-item').count() == 1)
    # exportar e importar en «otra PC»
    with pg.expect_download() as dl: pg.locator('.rmd-acu-barra [data-acc=exportar]').click()
    x = OUT / 'acuerdos.xlsx'; dl.value.save_as(str(x)); ok('Excel exportado', x.stat().st_size > 2000, x.stat().st_size)
    import openpyxl
    wb = openpyxl.load_workbook(str(x)); ok('hojas Acuerdos, Pasos, Historial e Imagenes', wb.sheetnames == ['Acuerdos', 'Pasos', 'Historial', 'Imagenes'], wb.sheetnames)
    copia = dict(kv); kv.clear()
    pg.locator('.rmd-acu-barra [data-acc=actualizar]').click(); pg.wait_for_timeout(700)
    ok('otra PC vacía', pg.locator('.rmd-acu-item').count() == 0 and 'Todavía no hay acuerdos' in pg.locator('.rmd-acu-lista').inner_text())
    with pg.expect_file_chooser() as fc: pg.locator('.rmd-acu-barra [data-acc=importar]').click()
    fc.value.set_files(str(x)); pg.wait_for_timeout(1500)
    pg.locator('.rmd-modal button', has_text='Importar').click(); pg.wait_for_timeout(1200)
    ok('importa los 2', len(kv) == 2, list(kv))
    nid = n.get('id'); imp = kv.get('acuerdo:' + str(nid), {})
    ok('importar conserva imágenes, pasos y casillas', [i['src'] for i in imp.get('imagenes', [])] == [i['src'] for i in n['imagenes']] and imp['pasos'][0]['rPor'] and imp['faseDesde'] == 2, imp.get('pasos'))
    ok('importar conserva el estado anulado', kv['acuerdo:sr1']['estado'] == 'anulado')
    # eliminar
    pg.locator('.rmd-acu-item', has_text='Despejes').locator('[data-acc=borrar]').click(); pg.wait_for_timeout(300)
    pg.locator('.rmd-modal button', has_text='Eliminar').click(); pg.wait_for_timeout(700)
    ok('eliminar', 'acuerdo:' + str(nid) not in kv and pg.locator('.rmd-acu-item').count() == 0)
    # cambiar de pestaña
    pg.locator('[role=tab]', has_text='Paso').first.click(); pg.wait_for_timeout(700)
    ok('Paso se ve al volver y Acuerdos no', pg.evaluate("(" + VIS + ")('Paso')") and not pg.locator('.rmd-acu-tab').is_visible())
    pg.locator('[role=tab]', has_text='Acuerdos').click(); pg.wait_for_timeout(500)
    pg.locator('[role=tab]', has_text='Controles de Cambios').click(); pg.wait_for_timeout(700)
    ok('de Acuerdos a Controles de Cambios: solo CC', pg.locator('.rmd-cc-pestana').is_visible() and not pg.locator('.rmd-acu-tab').is_visible() and not pg.evaluate("(" + VIS + ")('Paso')"))
    pg.locator('[role=tab]', has_text='Acuerdos').click(); pg.wait_for_timeout(700)
    ok('de CC a Acuerdos: solo Acuerdos', pg.locator('.rmd-acu-tab').is_visible() and not pg.locator('.rmd-cc-pestana').is_visible())
    dlg1 = pg.evaluate("(() => { const r = document.querySelector('.sapMDialog').getBoundingClientRect(); return [r.width, r.height]; })()")
    ok('ventana sin cambio de tamaño', dlg0 == dlg1, (dlg0, dlg1))
    # apagar la opción
    pg.evaluate("window.__rmdStats.panel.abrir(true)"); pg.wait_for_timeout(300)
    pg.locator("#rmd-ui-panel input[data-k=acuerdos]").click(force=True); pg.wait_for_timeout(900)
    tabs2 = pg.evaluate(TABS)
    ok('al apagar se quita la pestaña y vuelve una del portal', not any(t.endswith('Acuerdos') for t in tabs2) and pg.locator('.rmd-acu-tab').count() == 0 and pg.evaluate("(" + VIS + ")('Estructura')"), tabs2)
    pg.locator("#rmd-ui-panel input[data-k=acuerdos]").click(force=True); pg.wait_for_timeout(900)
    ok('al encender vuelve', pg.evaluate(TABS)[-1].endswith('Acuerdos'))
    pg.locator("#rmd-ui-panel input[data-k=controlescambio]").click(force=True); pg.wait_for_timeout(900)
    pg.locator("#rmd-ui-panel input[data-k=controlescambio]").click(force=True); pg.wait_for_timeout(900)
    t4 = pg.evaluate(TABS); ok('al apagar y encender Controles de Cambios queda a la izquierda de Acuerdos', t4[-2].endswith('Controles de Cambios') and t4[-1].endswith('Acuerdos'), t4)
    ok('un separador antes de cada pestaña nuestra', pg.evaluate("sap.ui.getCore().byId('itbMaestra').getItems().slice(-4).map((x) => x.getMetadata().getName().split('.').pop()).join()") == 'IconTabSeparator,IconTabFilter,IconTabSeparator,IconTabFilter')
    ok('sin errores de página', not errs, errs[:3])
    b.close()
print('RESUMEN', sum(r[1] for r in res), '/', len(res))
