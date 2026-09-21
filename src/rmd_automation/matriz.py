"""Sugerencia previa al ingreso: ¿el producto figura como pendiente en la matriz de RMD priorizados?

Hoja "PLANTA 2 DOC TEC" del libro "Matriz RMD priorizados v1 KZ 2026.xlsx" (exportada a .xlsx o .csv).
El analista anota ahí, antes del ingreso, los RMD pendientes: columna G = ingreso, K = autorización
(valor "PENDIENTE") y O = "observación adicional" (el motivo). No hay código de RMD: se busca por
nombre de producto y etapa. Es solo una sugerencia; nunca bloquea el ingreso.
"""
from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path
from typing import List

from .snapshot import normalizar

COL_INGRESO, COL_AUTORIZACION, COL_OBSERVACION = 6, 10, 14  # G, K, O (base 0)


@dataclass
class Pendiente:
    fila: int
    producto: str
    etapa: str
    ingreso: str
    autorizacion: str
    motivo: str


def _filas(ruta: Path, hoja: str) -> List[List[str]]:
    if ruta.suffix.lower() == ".csv":
        with open(ruta, encoding="utf-8-sig", newline="") as f:
            return [[c.strip() for c in r] for r in csv.reader(f)]
    try:
        from openpyxl import load_workbook
    except ImportError as e:  # pragma: no cover
        raise RuntimeError("Para leer .xlsx instala openpyxl (pip install openpyxl) o exporta la hoja a .csv.") from e
    wb = load_workbook(ruta, read_only=True, data_only=True)
    nombre = next((n for n in wb.sheetnames if normalizar(n) == normalizar(hoja)), None)
    if nombre is None:
        raise ValueError(f"El libro no tiene la hoja {hoja!r}. Hojas: {wb.sheetnames}")
    return [["" if c is None else str(c).strip() for c in r] for r in wb[nombre].iter_rows(values_only=True)]


def _columna(cabecera: List[str], *claves: str) -> int | None:
    for i, c in enumerate(cabecera):
        if any(k in normalizar(c) for k in claves):
            return i
    return None


def _coincidencias(ruta, producto: str, etapa: str, hoja: str) -> List[Pendiente]:
    """Todas las filas de la matriz cuyo producto (y etapa, si se da) coinciden."""
    filas = _filas(Path(ruta), hoja)
    # La cabecera es la fila que tiene a la vez la columna del nombre y la de ETAPA (las filas 1-3 son leyendas).
    ini = next(
        (i for i, r in enumerate(filas) if _columna(r, "NOMBRE", "PRODUCTO") is not None and _columna(r, "ETAPA") is not None),
        0,
    )
    cab = filas[ini] if filas else []
    c_prod = _columna(cab, "NOMBRE", "PRODUCTO") or 0
    c_etapa = _columna(cab, "ETAPA")
    p, e = normalizar(producto), normalizar(etapa)
    out: List[Pendiente] = []
    for n, r in enumerate(filas[ini + 1:], ini + 2):
        r = r + [""] * (COL_OBSERVACION + 1 - len(r))
        if p not in normalizar(r[c_prod]):
            continue
        et = r[c_etapa] if c_etapa is not None and c_etapa < len(r) else ""
        if e and et and e not in normalizar(et) and normalizar(et) not in e:
            continue
        out.append(Pendiente(n, r[c_prod], et, r[COL_INGRESO], r[COL_AUTORIZACION], r[COL_OBSERVACION]))
    return out


def buscar_pendientes(ruta: str | Path, producto: str, etapa: str = "", hoja: str = "PLANTA 2 DOC TEC") -> List[Pendiente]:
    """Filas del producto cuyo ingreso o autorización dice PENDIENTE."""
    return [
        x for x in _coincidencias(ruta, producto, etapa, hoja)
        if "PENDIENTE" in normalizar(x.ingreso) or "PENDIENTE" in normalizar(x.autorizacion)
    ]


def revisar(ruta: str | Path, producto: str, etapa: str = "", hoja: str = "PLANTA 2 DOC TEC") -> List[str]:
    """Avisos para el usuario. Lista vacía = todo normal (no se muestra nada).

    - el producto/etapa no está en la matriz -> hay que incluirlo (no es habitual);
    - una fila del producto sin estado de ingreso -> hay que completarla;
    - una fila PENDIENTE -> se muestra el motivo (columna O).
    """
    filas = _coincidencias(ruta, producto, etapa, hoja)
    etiqueta = f"{producto}" + (f" [{etapa}]" if etapa else "")
    if not filas:
        return [f"{etiqueta} NO figura en la matriz de pendientes: debe incluirse (no es habitual)."]
    avisos: List[str] = []
    for x in filas:
        if not normalizar(x.ingreso):
            avisos.append(
                f"fila {x.fila} ({x.producto} [{x.etapa}]): sin estado de ingreso, ni PENDIENTE ni INGRESADO: debe completarse."
            )
        elif "PENDIENTE" in normalizar(x.ingreso) or "PENDIENTE" in normalizar(x.autorizacion):
            avisos.append(
                f"fila {x.fila} ({x.producto} [{x.etapa}]): ingreso={x.ingreso or '-'}, autorización={x.autorizacion or '-'} "
                f"— motivo: {x.motivo or '(sin observación)'}"
            )
    return avisos


def a_texto(avisos: List[str]) -> str:
    if not avisos:
        return ""
    return "Matriz de priorizados (aviso previo al ingreso):\n" + "\n".join(f"  - {a}" for a in avisos)
