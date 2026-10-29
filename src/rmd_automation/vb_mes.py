"""Excel de los RMD trabajados en un mes: última y penúltima versión con sus V°B° (solo lectura).

Lee el portal ya abierto en el Chromium de pruebas (CDP), con el modelo OData `mainModelv2` de la lista
«Configuración Manufactura Digital», igual que `maestro_portal`. No escribe nada en SAP.

- Un RMD es la cadena de versiones que comparte `codigoversionprincipal` (o su propio código si es la primera).
- Entra si alguna de sus versiones se registró, ingresó (fecha al inicio de la Observación) o autorizó en el mes.
  Las solicitudes sin código (Solicitado, Solicitud Aprobada/Rechazada) no cuentan: aún no tienen estructura.
- V°B° = pasos activos de la estructura PROCEDIMIENTO con tipo de dato «Realizado por y Visto bueno» (441).
  Aparte se informan los de tipo «Visto bueno» (440), que no entran en la comparación.
- Vigente = la versión Autorizada más alta.
"""
from __future__ import annotations

import collections
from datetime import date
from pathlib import Path
from typing import Dict, List

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from playwright.sync_api import sync_playwright

PLANTAS = {"1020": "ATE", "1021": "LIMA"}

_JS_MODELO = """
() => {
  const visible = (e) => !!(e && e.offsetParent);
  const btn = [...document.querySelectorAll('button')].find((b) => visible(b) && b.title === 'Exportar');
  const ctl = btn && sap.ui.getCore().byId(btn.id);
  const reg = ctl && ctl.mEventRegistry && ctl.mEventRegistry.press && ctl.mEventRegistry.press[0];
  const vista = reg && reg.oListener && reg.oListener.getView && reg.oListener.getView();
  window.__rmdVbModelo = vista && vista.getModel('mainModelv2');
  if (!window.__rmdVbModelo) throw new Error('abre la lista «Configuración Manufactura Digital» en el portal');
}
"""

_JS_COMUN = """
  const m = window.__rmdVbModelo, F = sap.ui.require('sap/ui/model/Filter');
  const leer = (ent, filters, params) => new Promise((ok, mal) => m.read('/' + ent, { filters, urlParameters: params,
    success: (r) => ok(r.results || []), error: (e) => mal(new Error(`${ent}: el portal no respondió (${(e && e.statusCode) || 'sin código'})`)) }));
  const todo = async (ent, filters, params, clave) => {
    const out = [];
    for (let i = 0; i < 100; i++) {
      const p = await leer(ent, filters, { ...params, $top: '1000', $skip: String(i * 1000), $orderby: clave });
      out.push(...p); if (p.length < 1000) break;
    }
    return out;
  };
"""

# Fecha local (Lima); si viene a medianoche UTC es una fecha sin hora y se toma tal cual.
_JS_MAESTRO = """
async () => {""" + _JS_COMUN + """
  const dd = (n) => String(n).padStart(2, '0');
  const f = (x) => (!(x instanceof Date) || isNaN(x) ? '' : x.getUTCHours() + x.getUTCMinutes() + x.getUTCSeconds() === 0
    ? x.toISOString().slice(0, 10) : `${x.getFullYear()}-${dd(x.getMonth() + 1)}-${dd(x.getDate())}`);
  const filas = await todo('MD', [], { $expand: 'estadoIdRmd,sucursalId',
    $select: 'mdId,codigo,version,descripcion,nivelTxt,areaRmdTxt,codDefectoReceta,fechaRegistro,fechaAutorizacion,observacion,codigoversionprincipal,estadoIdRmd/contenido,sucursalId/codigo' }, 'mdId');
  const est = await todo('ESTRUCTURA', [], { $select: 'estructuraId,descripcion' }, 'estructuraId');
  return { est: est.map((e) => ({ id: e.estructuraId, d: e.descripcion || '' })),
    filas: [...new Map(filas.map((x) => [x.mdId, x])).values()].map((x) => ({ md_id: x.mdId, codigo: x.codigo || '', version: x.version || 0,
      producto: x.descripcion || '', etapa: x.nivelTxt || '', seccion: x.areaRmdTxt || '', material: x.codDefectoReceta || '',
      registro: f(x.fechaRegistro), autorizacion: f(x.fechaAutorizacion), observacion: x.observacion || '',
      principal: x.codigoversionprincipal || '', estado: (x.estadoIdRmd && x.estadoIdRmd.contenido) || '',
      planta: (x.sucursalId && x.sucursalId.codigo) || '' })) };
}
"""

_JS_PASOS = """
async ([ids, procedimiento]) => {""" + _JS_COMUN + """
  const res = {};
  for (let i = 0; i < ids.length; i += 10) {
    const lote = ids.slice(i, i + 10), porMd = new F({ filters: lote.map((id) => new F('mdId_mdId', 'EQ', id)), and: false });
    const [es, ps] = await Promise.all([
      todo('MD_ESTRUCTURA', [porMd], { $select: 'mdEstructuraId,activo' }, 'mdEstructuraId'),
      todo('MD_ES_PASO', [porMd, new F({ filters: [new F('tipoDatoId_iMaestraId', 'EQ', 441), new F('tipoDatoId_iMaestraId', 'EQ', 440)], and: false })],
        { $select: 'mdEstructuraPasoId,mdId_mdId,mdEstructuraId_mdEstructuraId,estructuraId_estructuraId,activo,tipoDatoId_iMaestraId' }, 'mdEstructuraPasoId')]);
    const activas = new Set(es.filter((e) => e.activo !== false).map((e) => e.mdEstructuraId));
    lote.forEach((id) => { res[id] = { vb: 0, solo_vb: 0 }; });
    ps.filter((p) => p.activo !== false && activas.has(p.mdEstructuraId_mdEstructuraId) && procedimiento.includes(p.estructuraId_estructuraId))
      .forEach((p) => { res[p.mdId_mdId][+p.tipoDatoId_iMaestraId === 441 ? 'vb' : 'solo_vb'] += 1; });
  }
  return res;
}
"""


def _fecha(texto: str):
    return date.fromisoformat(texto) if texto else None


def _ingreso(observacion: str):
    t = observacion[:8]
    try:
        return date(int(t[:4]), int(t[4:6]), int(t[6:8])) if t.isdigit() else None
    except ValueError:
        return None


def _elegir(registros: List[dict]) -> dict:
    """Si un número de versión se repite, vale la no cancelada más reciente."""
    return max(registros, key=lambda x: (x["estado"] != "Cancelado", x["registro"]))


def armar_filas(maestro: List[dict], anio: int, mes: int) -> List[dict]:
    en_mes = lambda d: d is not None and d.year == anio and d.month == mes
    cadenas: Dict[str, List[dict]] = collections.defaultdict(list)
    for x in maestro:
        if x["codigo"] and x["planta"] in PLANTAS:
            cadenas[x["principal"] or x["codigo"]].append(x)
    filas = []
    for principal, regs in cadenas.items():
        movs = collections.defaultdict(list)
        for x in regs:
            for nombre, d in (("Registro", _fecha(x["registro"])), ("Ingreso", _ingreso(x["observacion"])),
                              ("Autorización", _fecha(x["autorizacion"]))):
                if en_mes(d):
                    movs[nombre].append(d)
        if not movs:
            continue
        por_version = collections.defaultdict(list)
        for x in regs:
            por_version[x["version"]].append(x)
        versiones = sorted(por_version)
        ult = _elegir(por_version[versiones[-1]])
        pen = _elegir(por_version[versiones[-2]]) if len(versiones) > 1 else None
        autorizadas = [v for v in versiones if any(x["estado"] == "Autorizado" for x in por_version[v])]
        filas.append({"principal": principal, "ult": ult, "pen": pen, "vigente": autorizadas[-1] if autorizadas else None,
                      "movimientos": ", ".join(n for n in ("Registro", "Ingreso", "Autorización") if n in movs),
                      "ultimo_mov": max(d for ds in movs.values() for d in ds)})
    return filas


def leer_portal(cdp_url: str, anio: int, mes: int) -> List[dict]:
    with sync_playwright() as playwright:
        browser = playwright.chromium.connect_over_cdp(cdp_url)
        frame = next((f for p in browser.contexts[0].pages for f in p.frames if "ui5appruntime" in f.url), None)
        if frame is None:
            raise RuntimeError("Abre el portal RMD (Configuración RMD) en ese Chromium e inicia sesión.")
        frame.evaluate(_JS_MODELO)
        datos = frame.evaluate(_JS_MAESTRO)
        filas = armar_filas(datos["filas"], anio, mes)
        procedimiento = [e["id"] for e in datos["est"] if e["d"].strip().upper().startswith("PROCEDIMIENTO")]
        ids = [r[k]["md_id"] for r in filas for k in ("ult", "pen") if r[k]]
        conteos = frame.evaluate(_JS_PASOS, [ids, procedimiento])
        # No se llama a browser.close(): cerraría el Chromium de la persona.
    for r in filas:
        for k in ("ult", "pen"):
            if r[k]:
                r[k] = {**r[k], **conteos[r[k]["md_id"]]}
    return filas


COLUMNAS = [  # (título, ancho, formato)
    ("Planta", 8, None), ("Etapa", 15, None), ("Producto", 44, None), ("Sección", 18, None), ("Material", 12, None),
    ("Movimiento en el mes", 22, None), ("Fecha último movimiento", 13, "dd/mm/yyyy"),
    ("Últ. versión", 8, None), ("Últ. código RMD", 13, None), ("Últ. estado", 13, None),
    ("Últ. F. registro", 12, "dd/mm/yyyy"), ("Últ. F. ingreso", 12, "dd/mm/yyyy"), ("Últ. F. autorización", 13, "dd/mm/yyyy"),
    ("Últ. V°B° (Realizado por y Visto bueno)", 14, None), ("Últ. solo «Visto bueno»", 10, None),
    ("Penúlt. versión", 9, None), ("Penúlt. código RMD", 13, None), ("Penúlt. estado", 13, None),
    ("Penúlt. V°B° (Realizado por y Visto bueno)", 14, None), ("Penúlt. solo «Visto bueno»", 10, None),
    ("Diferencia V°B°", 10, None), ("Variación", 14, None), ("Versión vigente", 9, None), ("¿Última es la vigente?", 11, None),
]


def _variacion(dif):
    if dif is None:
        return "Sin versión anterior"
    return "Aumentó" if dif > 0 else "Disminuyó" if dif < 0 else "Se mantuvo"


def escribir_excel(filas: List[dict], salida: Path, anio: int, mes: int) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "RMD del mes"
    ws.append([c[0] for c in COLUMNAS])
    orden = lambda r: (PLANTAS[r["ult"]["planta"]], r["ult"]["etapa"], r["ult"]["producto"])
    for r in sorted(filas, key=orden):
        u, p = r["ult"], r["pen"]
        vb_u = u["vb"]
        vb_p = p["vb"] if p else None
        dif = vb_u - vb_p if p else None
        ws.append([PLANTAS[u["planta"]], u["etapa"], u["producto"], u["seccion"], u["material"],
                   r["movimientos"], r["ultimo_mov"],
                   u["version"], u["codigo"], u["estado"], _fecha(u["registro"]), _ingreso(u["observacion"]), _fecha(u["autorizacion"]),
                   vb_u, u["solo_vb"],
                   p and p["version"], p and p["codigo"], p and p["estado"], vb_p, p and p["solo_vb"],
                   dif, _variacion(dif), r["vigente"], "Sí" if r["vigente"] == u["version"] else "No"])
    for i, (_, ancho, formato) in enumerate(COLUMNAS, 1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
        if formato:
            for (celda,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                celda.number_format = formato
    for celda in ws[1]:
        celda.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 32
    ultima = get_column_letter(len(COLUMNAS))
    tabla = Table(displayName="RMDMes", ref=f"A1:{ultima}{ws.max_row}")
    tabla.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tabla)
    ws.freeze_panes = "D2"
    col_var = get_column_letter(22)
    rango = f"{col_var}2:{col_var}{ws.max_row}"
    for texto, color in (("Aumentó", "C6EFCE"), ("Disminuyó", "FFC7CE"), ("Se mantuvo", "FFEB9C")):
        ws.conditional_formatting.add(rango, CellIsRule(operator="equal", formula=[f'"{texto}"'], fill=PatternFill("solid", fgColor=color)))

    res = wb.create_sheet("Resumen")
    res.append([f"RMD trabajados en {mes:02d}/{anio}: conteo por variación de V°B°"])
    res["A1"].font = Font(bold=True)
    variaciones = ["Aumentó", "Disminuyó", "Se mantuvo", "Sin versión anterior"]
    res.append(["Planta", "Etapa", *variaciones, "Total"])
    for c in res[2]:
        c.font = Font(bold=True)
    cuenta = collections.Counter()
    for r in filas:
        u, p = r["ult"], r["pen"]
        dif = u["vb"] - p["vb"] if p else None
        cuenta[(PLANTAS[u["planta"]], u["etapa"], _variacion(dif))] += 1
    for planta, etapa in sorted({k[:2] for k in cuenta}):
        n = [cuenta[(planta, etapa, v)] for v in variaciones]
        res.append([planta, etapa, *n, sum(n)])
    n = [sum(c for k, c in cuenta.items() if k[2] == v) for v in variaciones]
    res.append(["Total", "", *n, sum(n)])
    for c in res[res.max_row]:
        c.font = Font(bold=True)
    res.append([])
    for nota in (
        "Criterios",
        "• Un RMD es la cadena de versiones de un mismo código principal; entra si alguna versión se registró, ingresó "
        "(fecha al inicio de la Observación) o autorizó en el mes. Solo Planta Ate y Planta Lima.",
        "• Las solicitudes sin código (Solicitado, Solicitud Aprobada o Rechazada) no cuentan como versión: aún no tienen estructura.",
        "• Si un número de versión se repite, se toma la no cancelada más reciente.",
        "• V°B° = pasos activos de la estructura PROCEDIMIENTO con tipo de dato «Realizado por y Visto bueno». "
        "Los de tipo «Visto bueno» solo se informan aparte y no entran en la diferencia.",
        "• Diferencia = V°B° de la última versión − V°B° de la penúltima.",
        "• Versión vigente = la versión Autorizada más alta. Ojo: al suspender, SAP cambia la fecha de autorización por la de la suspensión.",
        f"• Leído del portal (solo lectura) el {date.today():%d/%m/%Y}.",
    ):
        res.append([nota])
    res.column_dimensions["A"].width = 12
    res.column_dimensions["B"].width = 16
    for col in "CDEFG":
        res.column_dimensions[col].width = 12
    salida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(salida)
