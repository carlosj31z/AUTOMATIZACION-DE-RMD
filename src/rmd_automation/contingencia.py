"""Contingencia de RMD: ¿la carpeta de SharePoint tiene el PDF de cada RMD vigente?

Carpeta `04.-RMD/3.-CONTINGENCIA`: PLANTA / LÍNEA / PRODUCTO / ETAPA / <pdf>, con las versiones viejas en
`OBSOLETO/`. Cada PDF es el master impreso desde el portal y su nombre empieza con el código del RMD y la
versión, en alguna de estas formas:

    2202403301-FAB-5-_GASEOVET_80mg_mL_SUS 1000 L.pdf     código-ETAPA-versión-…
    2202607732-5-GASEOVET_80mg_mL_SUS_….pdf               código-versión-… (nombre de la descarga)
    2202506416-ENV3-ELECTROLIGHT_….pdf                    código-ETAPAversión-…
    13407-ACO-1-_GASEOVET_….pdf                           código antiguo (pre-migración)

Vigente = estado «Autorizado» en el maestro (exportación SAPUI5 de la lista «Configuración Manufactura
Digital», la misma de los libros «BD RMD <MES> - P1-P2.xlsx»). Solo lee: no escribe en SharePoint.
"""
from __future__ import annotations

import csv
import re
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path, PurePosixPath
from typing import Dict, Iterable, List, Optional

from .snapshot import normalizar

ESTADO_VIGENTE = "AUTORIZADO"
ETAPAS = {"FAB": "FABRICACION", "ENV": "ENVASE", "ACO": "ACONDICIONADO", "REC": "RECUBRIMIENTO", "INS": "INSPECCION"}

# Se toleran espacios junto a los guiones y «_» tras la versión («2202607906- INS-1-…», «…-INS-2_…», «15443-AACO1-2-…»).
_NOMBRES = [
    re.compile(r"^(?P<codigo>\d{4,10})\s*-\s*(?:(?P<etapa>[A-Za-z]{3,4})\d?\s*-\s*)?(?P<version>\d{1,3})\s*[-_]"),
    re.compile(r"^(?P<codigo>\d{4,10})\s*-\s*(?P<etapa>[A-Za-z]{3})(?P<version>\d{1,3})\s*[-_]"),
]
_OBSOLETO = re.compile(r"OB?S?OLETOS?\s*\d*")  # también «OSOLETO», «OBOLETO», «OBSOLETO1»
_COPIA = re.compile(r"\s\(\d+\)\.pdf$", re.IGNORECASE)


@dataclass
class Rmd:
    codigo: str
    version: str
    estado: str
    descripcion: str
    etapa: str
    planta: str = ""
    seccion: str = ""
    agrupador: str = ""
    material: str = ""
    fecha_autorizacion: str = ""

    @property
    def vigente(self) -> bool:
        return normalizar(self.estado) == ESTADO_VIGENTE


@dataclass
class Pdf:
    ruta: str  # relativa a 3.-CONTINGENCIA, con «/»
    codigo: str
    version: str
    etapa: str  # abreviatura del nombre (FAB/ENV/ACO…), vacía si el nombre no la trae
    obsoleto: bool
    copia: bool  # «… (1).pdf»: otra descarga del mismo archivo

    @property
    def carpeta(self) -> str:
        """Carpeta de la etapa (sin OBSOLETO), donde debería estar el vigente."""
        partes = PurePosixPath(self.ruta).parent.parts
        if partes and es_obsoleto(partes[-1]):
            partes = partes[:-1]
        return "/".join(partes)


@dataclass
class Brechas:
    faltantes: List[dict] = field(default_factory=list)
    desactualizados: List[dict] = field(default_factory=list)
    no_vigentes: List[dict] = field(default_factory=list)
    duplicados: List[dict] = field(default_factory=list)
    sin_codigo: List[dict] = field(default_factory=list)
    al_dia: int = 0
    vigentes: int = 0
    pdfs: int = 0


def es_obsoleto(carpeta: str) -> bool:
    return bool(_OBSOLETO.fullmatch(normalizar(carpeta)))


# ---------------------------------------------------------------- maestro

def _filas(ruta: Path) -> List[List[str]]:
    if ruta.suffix.lower() == ".csv":
        with open(ruta, encoding="utf-8-sig", newline="") as f:
            return [[c.strip() for c in r] for r in csv.reader(f)]
    try:
        from openpyxl import load_workbook
    except ImportError as e:  # pragma: no cover
        raise RuntimeError("Para leer .xlsx instala openpyxl (pip install openpyxl) o exporta la hoja a .csv.") from e
    wb = load_workbook(ruta, read_only=True, data_only=True)
    # La exportación del portal es la primera hoja; los libros «BD RMD» le agregan hojas de resumen detrás.
    hoja = next((wb[n] for n in wb.sheetnames if "EXPORT" in normalizar(n)), wb[wb.sheetnames[0]])
    return [["" if c is None else str(c).strip() for c in r] for r in hoja.iter_rows(values_only=True)]


def _col(cab: List[str], nombre: str) -> Optional[int]:
    n = normalizar(nombre)
    return next((i for i, c in enumerate(cab) if normalizar(c) == n), None)


def leer_maestro(ruta) -> List[Rmd]:
    """Lee la exportación SAPUI5 del maestro (.xlsx o .csv). La cabecera es la fila con «Código» y «Estado»."""
    filas = _filas(Path(ruta))
    ini = next((i for i, r in enumerate(filas) if _col(r, "Código") is not None and _col(r, "Estado") is not None), None)
    if ini is None:
        raise ValueError(f"{ruta}: no encuentro la cabecera con las columnas «Código» y «Estado».")
    cab = filas[ini]
    c = {k: _col(cab, k) for k in ("Código", "Versión", "Estado", "Descripción", "Etapa", "Planta", "Sección",
                                   "Código Agrupador", "Código por Defecto", "Fecha Autorización")}

    def val(r: List[str], k: str) -> str:
        i = c[k]
        v = r[i] if i is not None and i < len(r) else ""
        return v[:-2] if re.fullmatch(r"\d+\.0", v) else v  # números que openpyxl devuelve como float

    out = []
    for r in filas[ini + 1:]:
        codigo = val(r, "Código")
        if not codigo.isdigit():  # solicitudes sin RMD creado, filas en blanco
            continue
        out.append(Rmd(codigo=codigo, version=val(r, "Versión"), estado=val(r, "Estado"),
                       descripcion=val(r, "Descripción"), etapa=val(r, "Etapa"), planta=val(r, "Planta"),
                       seccion=val(r, "Sección"), agrupador=val(r, "Código Agrupador"),
                       material=val(r, "Código por Defecto"), fecha_autorizacion=val(r, "Fecha Autorización")[:10]))
    return out


# ---------------------------------------------------------------- PDFs

def interpretar(ruta: str) -> Optional[Pdf]:
    """Código, versión y etapa a partir del nombre; None si el nombre no sigue ninguna convención."""
    ruta = ruta.replace("\\", "/").strip("/")
    p = PurePosixPath(ruta)
    m = next((m for r in _NOMBRES for m in [r.match(p.name)] if m), None)
    if not m:
        return None
    return Pdf(ruta=ruta, codigo=m["codigo"], version=m["version"], etapa=(m["etapa"] or "").upper()[-3:],
               obsoleto=any(es_obsoleto(x) for x in p.parts[:-1]), copia=bool(_COPIA.search(p.name)))


def listar_carpeta(raiz) -> List[str]:
    """PDF de una copia local de 3.-CONTINGENCIA (p. ej. sincronizada con OneDrive), relativos a la raíz."""
    raiz = Path(raiz)
    return sorted(str(p.relative_to(raiz)).replace("\\", "/") for p in raiz.rglob("*") if p.suffix.lower() == ".pdf")


def leer_listado(ruta) -> List[str]:
    """Un archivo de texto con una ruta por línea (relativa a 3.-CONTINGENCIA)."""
    with open(ruta, encoding="utf-8-sig") as f:
        return [x.strip() for x in f if x.strip().lower().endswith(".pdf")]


# ---------------------------------------------------------------- comparación

def _fila(r: Rmd) -> dict:
    return {"Código": r.codigo, "Versión": r.version, "Estado": r.estado, "Descripción": r.descripcion,
            "Etapa": r.etapa, "Planta": r.planta, "Sección": r.seccion, "Fecha autorización": r.fecha_autorizacion}


def comparar(maestro: Iterable[Rmd], rutas: Iterable[str]) -> Brechas:
    maestro = list(maestro)
    por_codigo: Dict[str, Rmd] = {r.codigo: r for r in maestro}
    # Las versiones de un mismo RMD comparten agrupador (o, si falta, material) y etapa.
    familia = lambda r: (r.agrupador or r.material, normalizar(r.etapa))  # noqa: E731
    vigente_de: Dict[tuple, List[Rmd]] = defaultdict(list)
    for r in maestro:
        if r.vigente and (r.agrupador or r.material):
            vigente_de[familia(r)].append(r)

    b = Brechas()
    pdfs: List[Pdf] = []
    for ruta in rutas:
        b.pdfs += 1
        p = interpretar(ruta)
        if p is None:
            if not any(es_obsoleto(x) for x in PurePosixPath(ruta).parts[:-1]):
                b.sin_codigo.append({"Archivo": ruta})
            continue
        pdfs.append(p)

    activos: Dict[str, List[Pdf]] = defaultdict(list)  # fuera de OBSOLETO
    todos: Dict[str, List[Pdf]] = defaultdict(list)
    for p in pdfs:
        todos[p.codigo].append(p)
        if not p.obsoleto:
            activos[p.codigo].append(p)

    # PDF vigentes repetidos (otra descarga del mismo código en la misma carpeta o en otra).
    for codigo, lista in sorted(activos.items()):
        if len(lista) > 1:
            for p in lista:
                b.duplicados.append({"Código": codigo, "Versión": p.version, "Archivo": p.ruta,
                                     "Copia «(n)»": "SÍ" if p.copia else ""})

    # PDF fuera de OBSOLETO cuyo RMD ya no está Autorizado.
    reemplazos: Dict[str, Pdf] = {}
    for codigo, lista in sorted(activos.items()):
        r = por_codigo.get(codigo)
        if r is not None and r.vigente:
            continue
        nuevo = vigente_de.get(familia(r), []) if r is not None else []
        sin_pdf = [n for n in nuevo if not activos.get(n.codigo)]
        for p in lista:
            fila = {"Archivo": p.ruta, "Código": codigo, "Versión PDF": p.version,
                    "Estado en SAP": r.estado if r else "No está en el maestro",
                    "Descripción": r.descripcion if r else ""}
            if sin_pdf:  # la versión vigente no tiene PDF: hay que generarlo y pasar este a OBSOLETO
                fila["Vigente: código"] = ", ".join(n.codigo for n in sin_pdf)
                fila["Vigente: versión"] = ", ".join(n.version for n in sin_pdf)
                for n in sin_pdf:
                    reemplazos.setdefault(n.codigo, p)
                b.desactualizados.append(fila)
            else:  # basta con pasarlo a OBSOLETO
                if nuevo:
                    fila["Vigente ya en la carpeta"] = ", ".join(n.codigo for n in nuevo)
                b.no_vigentes.append(fila)

    # RMD Autorizados sin PDF vigente.
    for r in maestro:
        if not r.vigente:
            continue
        b.vigentes += 1
        if activos.get(r.codigo):
            b.al_dia += 1
            continue
        if r.codigo in reemplazos:  # ya figura en «desactualizados»
            continue
        fila = _fila(r)
        # Dónde guardarlo: junto a otra versión del mismo RMD si alguna ya está en la carpeta.
        hermano = next((p for o in maestro if familia(o) == familia(r) and (r.agrupador or r.material)
                        for p in todos.get(o.codigo, [])), None)
        fila["Carpeta sugerida"] = hermano.carpeta if hermano else ""
        fila["Nota"] = "Solo está en OBSOLETO" if todos.get(r.codigo) else ""
        b.faltantes.append(fila)
    return b


# ---------------------------------------------------------------- salida

def resumen(b: Brechas) -> str:
    return "\n".join([
        f"RMD vigentes (Autorizados): {b.vigentes}",
        f"  con PDF de contingencia al día: {b.al_dia}",
        f"  con PDF solo de una versión anterior (desactualizados): {b.vigentes - b.al_dia - len(b.faltantes)}",
        f"  sin ningún PDF: {len(b.faltantes)}",
        f"PDF revisados: {b.pdfs}",
        f"  de RMD no vigentes fuera de OBSOLETO (pasarlos a OBSOLETO): {len(b.no_vigentes) + len(b.desactualizados)}",
        f"  repetidos: {len(b.duplicados)}",
        f"  sin código reconocible en el nombre: {len(b.sin_codigo)}",
    ])


def escribir_excel(b: Brechas, salida) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    for linea in resumen(b).splitlines():
        ws.append([linea])
    hojas = [("Sin PDF", b.faltantes), ("Desactualizados", b.desactualizados), ("Pasar a OBSOLETO", b.no_vigentes),
             ("Repetidos", b.duplicados), ("Nombre sin código", b.sin_codigo)]
    for titulo, filas in hojas:
        h = wb.create_sheet(titulo)
        if not filas:
            h.append(["(nada)"])
            continue
        cab = list(dict.fromkeys(k for f in filas for k in f))
        h.append(cab)
        for c in h[1]:
            c.font = Font(bold=True)
        for f in filas:
            h.append([f.get(k, "") for k in cab])
        h.freeze_panes = "A2"
        h.auto_filter.ref = h.dimensions
        for i, k in enumerate(cab, 1):
            ancho = max(len(str(k)), *(len(str(f.get(k, ""))) for f in filas[:500]))
            h.column_dimensions[get_column_letter(i)].width = min(max(ancho + 2, 8), 80)
    wb.save(salida)
